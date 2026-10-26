"""
ONE script, no MongoDB:  ASIN -> Amazon.in page (Playwright/Chromium) -> breadcrumb + details -> HSN -> your Excel.

  py asin_to_hsn.py SHEET.xlsx --limit=5        test: fetch 5 ASINs from Amazon, then build the output
  py asin_to_hsn.py SHEET.xlsx                  fetch every ASIN not fetched yet (resumes), then build the output
  py asin_to_hsn.py SHEET.xlsx --build-only     no fetching: just rebuild the Excel from what is saved
  py asin_to_hsn.py SHEET.xlsx --retry-failed   try ASINs that failed before again

Fetched details are saved per ASIN in  json/<ASIN>.json  (so you can stop with Ctrl+C and continue any time).
Output  <SHEET>_HSN_DONE.xlsx  - a clean new sheet:  STORE | Seller Gstin | Item Description | Asin | Breadcrumb | HSN Chapter Code |
   HSN Description | HSN Tariff Code | GST Rate | Confidence | Remark | Fetch status  (+ 'Your existing HSN' if your sheet had values)
ASINs not fetched yet are still classified from the Item Description alone (low confidence, flagged).
Setup once:  py -m pip install playwright openpyxl   &   py -m playwright install chromium
Amazon restricts automated page reading: use with your employer's approval; keep the delays.
"""
import json, os, random, re, sys, time
from collections import Counter

DELAY = (float(os.environ.get("BC_MIN", 4)), float(os.environ.get("BC_MAX", 8)))
OUT_DIR = "json"

def find_header(ws):
    for ri, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True), 1):
        h = [str(x).strip().lower() if x else "" for x in row]
        if "asin" in h and any(x.startswith("item desc") for x in h):
            return ri, {n: i + 1 for i, n in enumerate(h) if n}
    sys.exit("No header row with 'Item Description' and 'Asin' found")

def read_sheet(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=True).active
    hr, col = find_header(ws); ti = col[next(k for k in col if k.startswith("item desc"))] - 1; ai = col["asin"] - 1
    items = {}
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        a = str(row[ai]).strip().upper() if row[ai] else ""
        if a: items.setdefault(a, str(row[ti] or ""))
    return items

# ------------------------------------------------------------------ fetching
def fetch(items, args):
    import asin_scraper as S
    todo = []
    for a in items:
        if os.path.exists(f"{OUT_DIR}/{a}.json"): continue
        if os.path.exists(f"{OUT_DIR}/{a}.failed.json") and "--retry-failed" not in args: continue
        todo.append(a)
    limit = next((int(x.split("=")[1]) for x in args if x.startswith("--limit=")), None)
    if limit: todo = todo[:limit]
    print(f"{len(items)} unique ASINs | {len(todo)} to fetch in this run")
    if not todo: return
    os.makedirs(OUT_DIR, exist_ok=True)
    from playwright.sync_api import sync_playwright
    fails = 0
    with sync_playwright() as p:
        ctx = p.chromium.launch_persistent_context("browser_profile", headless=bool(os.environ.get("BC_HEADLESS")),
                                                   locale="en-IN", viewport={"width": 1280, "height": 900})
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        try:
            for i, a in enumerate(todo, 1):
                try:
                    page.goto(S.BASE.format(asin=a), wait_until="domcontentloaded", timeout=45000)
                    try: page.wait_for_selector("#productTitle, #captchacharacters, #g", timeout=12000)
                    except Exception: pass
                    d = S.parse_in_page(page, a)
                    while d["captcha"]:
                        input("  CAPTCHA - solve it in the browser window, then press Enter here... ")
                        page.wait_for_timeout(1500); d = S.parse_in_page(page, a)
                    if d["title"]:
                        d["status"] = "ok"
                        json.dump(d, open(f"{OUT_DIR}/{a}.json", "w", encoding="utf-8"), ensure_ascii=False, indent=2)
                        fails = 0; msg = d["breadcrumb"][:70] or "(no breadcrumb on page)"
                    else:
                        json.dump({"asin": a, "status": "empty/not found"}, open(f"{OUT_DIR}/{a}.failed.json", "w"))
                        fails += 1; msg = "empty / not found"
                except Exception as e:
                    fails += 1; msg = f"error {e.__class__.__name__}"
                print(f"[{i}/{len(todo)}] {a} | {msg}")
                if fails >= 10:
                    print("10 failures in a row - stopping (blocked?). Wait a while, then run again; it continues."); break
                time.sleep(random.uniform(*DELAY))
        except KeyboardInterrupt:
            print("\nStopped - progress is saved.")
        ctx.close()

# ------------------------------------------------------------------ output
def build(src, items):
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment
    import build_hsn_records as B, asin_scraper as S
    from hsn_classify import Classifier
    wb = openpyxl.load_workbook(src, read_only=True); ws = wb.active; hr, col = find_header(ws)
    def cidx(*names):
        for n in names:
            for k, c in col.items():
                if k.strip() == n: return c - 1
        return None
    iStore, iGst = cidx("store"), cidx("seller gstin"); iT = next(c - 1 for k, c in col.items() if k.startswith("item desc")); iA = col["asin"] - 1
    iHsn = cidx("hsn"); iRate = cidx("rate")
    src_rows = [r for r in ws.iter_rows(min_row=hr + 1, values_only=True) if r[iA]]
    clf = Classifier(); res = {}
    llm = {}
    if os.path.exists("llm_results.jsonl"):                  # AI step results (llm_hsn.py) override the keyword/rule result
        for line in open("llm_results.jsonl", encoding="utf-8"):
            try: d = json.loads(line); llm[d["asin"]] = d
            except Exception: pass
        print(f"AI results found for {len(llm)} products")
    for a, title in items.items():
        p = None
        if os.path.exists(f"{OUT_DIR}/{a}.json"):
            try: p = B.from_master(S.to_master(json.load(open(f"{OUT_DIR}/{a}.json", encoding="utf-8")), title)); p["_f"] = "Fetched from Amazon"
            except Exception: p = None
        if not p:
            p = dict(title=title, brand="", product_type=None, browse_node_id="", leaf_category="", breadcrumb="", breadcrumb_source="none",
                     bullets=[], description="", features={}, data_source="")
            p["_f"] = "Failed on Amazon - description only" if os.path.exists(f"{OUT_DIR}/{a}.failed.json") else "Not fetched yet - description only"
        if not p["title"]: p["title"] = title
        blk = B.hsn_block(clf, p)
        if a in llm:
            d = llm[a]; conf = "AI-High" if d.get("confidence") == "high" else "AI-Review"
            nb = B.block_from_code(clf, d["code"], conf, ["AI choice: " + d.get("reason", "")] + (["AI not sure"] if d.get("confidence") == "low" else []), "AI")
            if nb: blk = nb
        res[a] = (blk, p)
    out_wb = openpyxl.Workbook(); o = out_wb.active; o.title = "HSN"
    heads = ["STORE", "Seller Gstin", "Item Description", "Asin", "Breadcrumb", "HSN Chapter Code", "HSN Description", "HSN Tariff Code",
             "GST Rate", "Confidence", "Remark", "Fetch status"]
    has_existing = iHsn is not None and any(r[iHsn] not in (None, "") for r in src_rows)
    if has_existing: heads.append("Your existing HSN")
    o.append(heads)
    for c in o[1]: c.font = Font(bold=True)
    fills = {"High": "C6EFCE", "AI-High": "A9D08E", "AI-Review": "DDEBF7", "Review": "FFEB9C", "Unreviewed": "F8CBAD", "Unclassified": "D9D9D9"}
    stat = Counter(); fs = Counter(); srcs = Counter()
    for r in src_rows:
        a = str(r[iA]).strip().upper(); h, p = res[a]
        conf = h["confidence"] if h else "Unclassified"; flags = list(h["flags"]) if h else ["Could not classify - manual"]
        if p["_f"] != "Fetched from Amazon" and not conf.startswith("AI"):
            conf = conf if conf == "Unclassified" else "Unreviewed"; flags.insert(0, "Item Description only - no breadcrumb")
        rate = ""
        if h: rate = f"{h['gst_percent']:g}%" if h["gst_percent"] is not None else h["gst_rate"]
        row = [r[iStore] if iStore is not None else "", r[iGst] if iGst is not None else "", r[iT], a, p["breadcrumb"],
               h["chapter_2"] if h else "", h["hsn_description"] if h else "", int(h["hsn_8"]) if h else "", rate, conf,
               "; ".join(flags), p["_f"]]
        if has_existing: row.append(r[iHsn] if r[iHsn] not in (None, "") else "")
        o.append(row); o.cell(o.max_row, 10).fill = PatternFill("solid", fgColor=fills[conf])
        stat[conf] += 1; fs[p["_f"]] += 1; srcs[(h or {}).get("classified_by", "-").split(" '")[0]] += 1
    for letter, w in zip("ABCDEFGHIJKLM", (16, 18, 55, 13, 50, 10, 55, 13, 9, 12, 60, 26, 14)): o.column_dimensions[letter].width = w
    o.freeze_panes = "A2"; o.auto_filter.ref = o.dimensions
    out = os.path.splitext(src)[0] + "_HSN_DONE.xlsx"
    try: out_wb.save(out)
    except PermissionError:                          # the file is open in Excel -> save under a new name instead
        out = os.path.splitext(out)[0] + time.strftime("_%H%M%S") + ".xlsx"
        print("(The usual output file is open in Excel - close it next time.)"); out_wb.save(out)
    print("Confidence:", dict(stat)); print("Rows by data:", dict(fs)); print("Classified by:", dict(srcs)); print("Saved ->", out)

def main():
    args = sys.argv[1:]
    if not args or args[0].startswith("--"): print(__doc__); return
    src = args[0]
    if not os.path.exists(src):                      # tolerate a hidden double extension (name.xlsx.xlsx) or a missing one
        for cand in (src + ".xlsx", os.path.splitext(src)[0]):
            if os.path.exists(cand): print(f"Using {cand}"); src = cand; break
        else: sys.exit(f"File not found: {src}\nFiles here: {', '.join(f for f in os.listdir('.') if f.lower().endswith('.xlsx'))}")
    items = read_sheet(src)
    if "--build-only" not in args: fetch(items, args)
    print("Building the Excel ..."); build(src, items)

if __name__ == "__main__":
    main()
