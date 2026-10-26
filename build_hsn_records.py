"""
ASIN -> important product features + breadcrumb -> HSN (2 / 4 / 6 / 8 digit) + descriptions + GST rate -> JSON + MongoDB

Product data is taken from (in this order):
   0) MongoDB  hsn_work.asin_master            (made by consolidate_asin.py - best source; HSN result is also written back into it)
   1) MongoDB  spapi_catalog.amazon.ins      (SP-API document)
   2) MongoDB  asin_scrape.products / json/<ASIN>.json   (details fetched from the Amazon.in page)
   3) the Item Description in your Excel      (title-only, low confidence)
Breadcrumb: from node_paths.csv (browse node id -> full path), or the scraped breadcrumb, or just the leaf name.

  py build_hsn_records.py ASIN_Missing.xlsx            every ASIN in the file
  py build_hsn_records.py --asin B0009MZJ6O            one ASIN
  py build_hsn_records.py file.xlsx --limit=200        first 200 (test)
  py build_hsn_records.py file.xlsx --no-db            JSON/Excel only

Output:  MongoDB  hsn_results.asin_hsn  (upsert by asin) | hsn_json/<ASIN>.json | <file>_HSN_RESULT.xlsx
Needs in the same folder: HSN json files, category_rules.json, hsn_chapter.py, hsn_code.py, hsn_classify.py
"""
import csv, json, os, re, sys
from datetime import datetime, timezone

import cfg
MONGO_URI = cfg.mongo_uri()
SRC_DB, SRC_COLL = "spapi_catalog", "amazon.ins"
SCRAPE_DB, SCRAPE_COLL = "asin_scrape", "products"
OUT_DB, OUT_COLL = "hsn_results", "asin_hsn"
MASTER_DB, MASTER_COLL = "hsn_work", "asin_master"      # made by consolidate_asin.py (preferred source)
NODE_PATHS = "node_paths.csv"
MARKET = "A21TJRUUN4KGV"

# attributes that help decide the HSN code (stored under "features")
KEY_ATTRS = ["item_name", "brand", "manufacturer", "material", "fabric_type", "item_type_keyword", "color", "size",
             "style", "model_number", "part_number", "unspsc_code", "number_of_items", "item_package_quantity",
             "power_source_type", "voltage", "wattage", "number_of_teeth", "capacity", "pattern", "target_audience_keyword",
             "department", "special_feature", "specific_uses_for_product", "included_components", "age_range_description",
             "recommended_uses_for_product", "ingredients", "item_form", "scent", "flavor", "unit_count", "composition"]
SKIP_ATTRS = {"externally_assigned_product_identifier", "item_dimensions", "item_package_dimensions", "item_package_weight",
              "item_weight", "list_price", "product_site_launch_date", "gpsr_safety_attestation", "batteries_required",
              "batteries_included", "recommended_browse_nodes", "supplier_declared_dg_hz_regulation", "country_of_origin",
              "merchant_suggested_asin", "purchasable_offer", "fulfillment_availability", "main_product_image_locator"}

def _val(v):
    """attribute value -> plain text"""
    if isinstance(v, dict):
        if "value" in v: return _val(v["value"])
        return " ".join(f"{k} {_val(x)}" for k, x in v.items() if k not in ("marketplace_id", "language_tag"))
    return str(v)

def attr_text(attrs, name):
    items = attrs.get(name) or []
    out = []
    for it in items:
        if isinstance(it, dict) and it.get("marketplace_id", MARKET) not in (MARKET, None): continue
        t = _val(it).strip() if it is not None else ""
        if t and t not in out: out.append(t)
    return "; ".join(out)

def load_node_paths(path=NODE_PATHS):
    m = {}
    if os.path.exists(path):
        for r in csv.DictReader(open(path, encoding="utf-8-sig")):
            if r.get("path"): m[str(r["node_id"]).strip()] = r["path"].strip()
    return m

def from_spapi(doc, node_paths):
    a = doc.get("attributes") or {}
    summ = next((s for s in doc.get("summaries", []) if s.get("marketplaceId") == MARKET), (doc.get("summaries") or [{}])[0])
    bc = summ.get("browseClassification") or {}
    node, leaf = str(bc.get("classificationId") or ""), bc.get("displayName") or ""
    pt = next((p.get("productType") for p in doc.get("productTypes", []) if p.get("marketplaceId") == MARKET),
              (doc.get("productTypes") or [{}])[0].get("productType"))
    feats = {k: attr_text(a, k) for k in KEY_ATTRS if attr_text(a, k)}
    for k in a:                                   # any other simple text attribute
        if k not in feats and k not in SKIP_ATTRS and k not in KEY_ATTRS:
            t = attr_text(a, k)
            if t and len(t) < 200: feats[k] = t
    bullets = [_val(b).strip() for b in (a.get("bullet_point") or [])]
    title = feats.get("item_name") or summ.get("itemName") or ""
    crumb = node_paths.get(node) or leaf
    return dict(title=title, brand=summ.get("brand") or feats.get("brand", ""), product_type=pt, browse_node_id=node,
                leaf_category=leaf, breadcrumb=crumb, breadcrumb_source="node_paths.csv" if node in node_paths else
                ("leaf only (breadcrumb not fetched yet)" if leaf else "none"),
                bullets=bullets, description=attr_text(a, "product_description"), features=feats, data_source="SP-API catalog (MongoDB)")

def from_scraped(d):
    specs = d.get("specifications") or {}
    return dict(title=d.get("title", ""), brand=d.get("brand", ""), product_type=None, browse_node_id=(d.get("breadcrumb_node_ids") or [""])[-1] or "",
                leaf_category=(d.get("breadcrumbs") or [""])[-1], breadcrumb=d.get("breadcrumb", ""),
                breadcrumb_source="Amazon page", bullets=d.get("bullets") or [], description=d.get("description", ""),
                features=specs, data_source="Amazon.in page (scraped)")

def from_master(d):
    return dict(title=d.get("title", ""), brand=d.get("brand", ""), product_type=d.get("product_type"),
                browse_node_id=d.get("classification_id", ""), leaf_category=d.get("classification_name", ""),
                breadcrumb=d.get("breadcrumb", ""), breadcrumb_source=d.get("breadcrumb_source", ""),
                bullets=d.get("bullets") or [], description=d.get("description", ""),
                features={**(d.get("features") or {}), **({"first_category": d["first_category"]} if d.get("first_category") else {})},
                data_source="asin_master (amazon.ins + amazon.in_formats)")

def hsn_block(clf, p):
    crumb = p["breadcrumb"]
    # leaf-only breadcrumb: borrow the parent from a reviewed rule if the leaf name is unique there
    if crumb and ">" not in crumb:
        hits = [k for k in clf.rules if k.endswith("> " + crumb.lower()) or k == crumb.lower()]
        heads = {clf.rules[k]["heading"] for k in hits}
        if len(hits) >= 1 and len(heads) == 1: crumb = clf.rules[hits[0]]["example"]
    feats = dict(p["features"])
    res = clf.classify(p["title"], crumb, p["bullets"], p["description"], feats)
    if not res: return None
    b = res["best"]; code = b["code"]
    if not hasattr(clf, "_gname"): clf._gname = {str(x["number"]): x.get("name", "") for x in clf.gst}
    gn = clf._gname
    return {"chapter_2": code[:2], "chapter_name": res["chapter_name"], "heading_4": res["heading"],
            "heading_description": gn.get(res["heading"]) or res["heading_description"], "subheading_6": code[:6],
            "subheading_description": gn.get(code[:6], ""), "hsn_8": code, "hsn_description": b["description"], "unit": b["unit"], "gst_rate": b["rate"],
            "gst_percent": _pct(b["rate"]), "confidence": res["confidence"], "classified_by": res["source"],
            "flags": res["flags"], "alternatives": [{"hsn_8": c["code"], "description": c["description"], "gst_rate": c["rate"], "score": c["score"]}
                                                    for c in res["candidates"][1:4]]}

def block_from_code(clf, code, conf, flags, by):
    """HSN block for a chosen 8-digit code (used by the AI step)."""
    import hsn_code as K
    if not hasattr(clf, "_bycode"): clf._bycode = clf.cidx["items"]
    if not hasattr(clf, "_gname"): clf._gname = {str(x["number"]): x.get("name", "") for x in clf.gst}
    it = clf._bycode.get(code)
    if not it: return None
    rates, label, hr = K.rate_for(clf.cidx, code); rate = K.fmt_rate(rates)
    ch = code[:2]
    return {"chapter_2": ch, "chapter_name": clf.chname.get(ch, ""), "heading_4": it["heading"],
            "heading_description": clf._gname.get(it["heading"]) or "", "subheading_6": code[:6],
            "subheading_description": clf._gname.get(code[:6], ""), "hsn_8": code, "hsn_description": it["description"],
            "unit": it.get("unit"), "gst_rate": rate, "gst_percent": _pct(rate), "confidence": conf, "classified_by": by,
            "flags": flags, "alternatives": []}

def _pct(rate):
    m = re.match(r"\s*([\d.]+)\s*%", rate or "")
    return float(m.group(1)) if m else None

def build_record(clf, asin, p, excel_title=""):
    rec = {"asin": asin, "title": p["title"] or excel_title, "brand": p["brand"], "product_type": p["product_type"],
           "browse_node_id": p["browse_node_id"], "leaf_category": p["leaf_category"], "breadcrumb": p["breadcrumb"],
           "breadcrumb_source": p["breadcrumb_source"], "bullets": p["bullets"][:8], "features": p["features"],
           "data_source": p["data_source"], "built_at": datetime.now(timezone.utc).isoformat()}
    rec["hsn"] = hsn_block(clf, {**p, "title": rec["title"]})
    return rec

def read_asins(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows[:30]) if "asin" in [str(x).strip().lower() if x else "" for x in r])
    h = [str(x).strip().lower() if x else "" for x in rows[hi]]
    ai = h.index("asin"); ti = next((i for i, x in enumerate(h) if x.startswith("item desc")), None)
    seen, out = set(), []
    for r in rows[hi + 1:]:
        a = str(r[ai]).strip().upper() if r[ai] else ""
        if a and a not in seen: seen.add(a); out.append((a, str(r[ti] or "") if ti is not None else ""))
    return out

def main():
    from hsn_classify import Classifier
    args = sys.argv[1:]
    if not args: print(__doc__); return
    use_db = "--no-db" not in args
    limit = next((int(a.split("=")[1]) for a in args if a.startswith("--limit=")), None)
    if "--asin" in args:
        items = [(args[args.index("--asin") + 1].upper(), "")]; base = "single"
    else:
        items = read_asins(args[0]); base = os.path.splitext(args[0])[0]
    if limit: items = items[:limit]
    print(f"{len(items)} ASIN(s) | loading HSN data ...")
    clf = Classifier(); node_paths = load_node_paths(); print(f"{len(node_paths)} fetched breadcrumb paths")
    src = scr = out = master = None
    if use_db:
        try:
            from pymongo import MongoClient
            cl = MongoClient(MONGO_URI, serverSelectionTimeoutMS=4000); cl.admin.command("ping")
            src, scr, out = cl[SRC_DB][SRC_COLL], cl[SCRAPE_DB][SCRAPE_COLL], cl[OUT_DB][OUT_COLL]
            master = cl[MASTER_DB][MASTER_COLL]
            out.create_index("asin", unique=True)
        except Exception as e: print(f"! MongoDB not reachable ({e.__class__.__name__}) - files only")
    os.makedirs("hsn_json", exist_ok=True); rows = []
    for i, (asin, xt) in enumerate(items, 1):
        p = None
        if master is not None:
            d = master.find_one({"asin": asin})
            if d: p = from_master(d)
        if src is not None and not p:
            d = src.find_one({"asin": asin})
            if d: p = from_spapi(d, node_paths)
            if not p:
                d = scr.find_one({"asin": asin}, {"_id": 0})
                if d: p = from_scraped(d)
        if not p and os.path.exists(f"json/{asin}.json"): p = from_scraped(json.load(open(f"json/{asin}.json", encoding="utf-8")))
        if not p:
            p = dict(title=xt, brand="", product_type=None, browse_node_id="", leaf_category="", breadcrumb="", breadcrumb_source="none",
                     bullets=[], description="", features={}, data_source="Excel description only")
        if not p["title"] and not xt: print(f"[{i}] {asin} no data"); continue
        rec = build_record(clf, asin, p, xt)
        json.dump(rec, open(f"hsn_json/{asin}.json", "w", encoding="utf-8"), ensure_ascii=False, indent=2)
        if out is not None: out.replace_one({"asin": asin}, rec, upsert=True)
        if master is not None and rec["hsn"] and master.count_documents({"asin": asin}, limit=1):
            master.update_one({"asin": asin}, {"$set": {"hsn": rec["hsn"], "hsn_built_at": rec["built_at"]}})
        h = rec["hsn"] or {}
        rows.append([asin, rec["title"], rec["breadcrumb"], rec["data_source"], h.get("chapter_2"), h.get("heading_4"), h.get("subheading_6"),
                     h.get("hsn_8"), h.get("hsn_description"), h.get("gst_rate"), h.get("confidence"), "; ".join(h.get("flags", []))])
        if i % 200 == 0: print(f"  {i}/{len(items)}")
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Asin", "Title", "Breadcrumb", "Data source", "Chapter (2)", "Heading (4)", "Subheading (6)", "HSN (8)", "HSN Description", "GST rate", "Confidence", "Remark"])
    for r in rows: ws.append(r)
    wb.save(base + "_HSN_RESULT.xlsx"); print("Done:", len(rows), "records ->", base + "_HSN_RESULT.xlsx")

if __name__ == "__main__":
    main()
